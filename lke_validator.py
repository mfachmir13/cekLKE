import pandas as pd
import requests
from openpyxl import load_workbook

def is_valid_google_drive_link(link):
    try:
        response = requests.head(link)
        return response.status_code == 200
    except:
        return False

def check_google_drive_folder_content(link):
    # Placeholder for Google Drive API logic
    pass

def validate_excel(file_path):
    wb = load_workbook(file_path)
    sheet = wb['jawaban']
    
    results = []
    
    for row in range(2, sheet.max_row + 1):  # Assuming there is a header row
        jawaban_cell = sheet.cell(row=row, column=11)  # Column K
        bukti_dukung_cell = sheet.cell(row=row, column=13)  # Column M
        note_cell = sheet.cell(row=row, column=18)  # Column R
        
        if jawaban_cell.fill.start_color == "0054A7FF":  # Check if the cell is blue
            if jawaban_cell.value is None or jawaban_cell.value == '':
                note_cell.value = "Jawaban is empty."
            else:
                note_cell.value = "Jawaban is present."
        
            if bukti_dukung_cell.value:
                if is_valid_google_drive_link(bukti_dukung_cell.value):
                    note_cell.value += " Google Drive link is accessible."
                    
                else:
                    note_cell.value += " Google Drive link is not accessible."
            else:
                note_cell.value += " No link provided."
        
        results.append(note_cell.value)
    
    wb.save("validated_results.xlsx")

if __name__ == "__main__":
    validate_excel("path_to_your_excel_file.xlsx")
